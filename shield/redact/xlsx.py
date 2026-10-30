from __future__ import annotations

import zipfile
from pathlib import Path

from openpyxl import load_workbook

from .docx import _apply, _targets


def _raw_zip_pass(path: Path, repl: dict[str, str]) -> None:
    with zipfile.ZipFile(path) as z:
        items = [(info, z.read(info.filename)) for info in z.infolist()]

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in items:
            if info.filename.endswith(".xml"):
                try:
                    text = data.decode("utf-8")
                except UnicodeDecodeError:
                    pass
                else:
                    data = _apply(text, repl).encode("utf-8")
            z.writestr(info, data)


def redact_xlsx(
    src: str | Path,
    entities,
    out: str | Path,
    aggressive: bool = False,
) -> Path:
    src = Path(src)
    out = Path(out)
    repl = _targets(entities)

    wb = load_workbook(src)
    for ws in wb.worksheets:
        try:
            ws._pivots = []
        except Exception:
            pass

        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = _apply(cell.value, repl)
                if cell.comment and cell.comment.text:
                    cell.comment.text = _apply(cell.comment.text, repl)

    for name in list(wb.defined_names.keys()):
        try:
            item = wb.defined_names[name]
            text = f"{name} {getattr(item, 'attr_text', '')}"
        except Exception:
            continue
        if any(v in text for v in repl):
            del wb.defined_names[name]

    wb.properties.creator = ""
    wb.properties.lastModifiedBy = ""

    tmp = out.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    _raw_zip_pass(tmp, repl)
    tmp.replace(out)

    return out

