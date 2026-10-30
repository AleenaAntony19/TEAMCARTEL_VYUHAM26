import json

import httpx
from docx import Document

from shield import llm
from shield.codescan import cognitive, scan
from shield.codescan.fix import plan_fix

DIAGNOSIS = "HIV positive, treated at St Marys ward 4"
TEXT = f"Patient note for the clinic team. Aadhaar 2341 2341 2346 on file. Diagnosis: {DIAGNOSIS}."
RAW_ID = "234123412346"


def reply_with(items, seen):
    def handler(request):
        seen.append(json.loads(request.content)["messages"][-1]["content"])
        return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps({"items": items})}}]})
    return httpx.MockTransport(handler)


def groq(monkeypatch):
    monkeypatch.setenv("SHIELD_LLM_PROVIDER", "groq")
    monkeypatch.setenv("GROQ_API_KEY", "test-key")
    monkeypatch.delenv("SHIELD_LOCAL_ONLY", raising=False)


def doc(tmp_path):
    d = Document()
    d.add_paragraph(TEXT)
    d.save(tmp_path / "note.docx")
    return tmp_path / "note.docx"


GOOD = {"value": DIAGNOSIS, "kind": "health data", "severity": "critical", "reason": f"{DIAGNOSIS} is health data"}
INVENTED = {"value": "not in the file at all", "kind": "X", "severity": "high", "reason": "made up"}


def test_known_values_blanked_invented_dropped_and_cache_is_masked(tmp_path, monkeypatch):
    groq(monkeypatch)
    path = doc(tmp_path)
    seen, stats = [], {}
    found = cognitive.scan_cognitive([path], tmp_path, True, False, stats, reply_with([GOOD, INVENTED], seen))
    assert RAW_ID not in seen[0].replace(" ", "") and DIAGNOSIS in seen[0]
    assert [f.severity for f in found] == ["medium"], "the model cannot claim critical, invented values are dropped"
    assert DIAGNOSIS not in found[0].rationale and found[0].rule == "AI_SENSITIVE"
    assert stats["cognitive"]["sent"] == 1
    cache = (tmp_path / cognitive.CACHE).read_text()
    assert DIAGNOSIS not in cache and RAW_ID not in cache
    assert (tmp_path / cognitive.CACHE).with_name(".gitignore").read_text() == "*\n"


def test_cached_findings_survive_without_network_until_file_changes(tmp_path, monkeypatch):
    groq(monkeypatch)
    path = doc(tmp_path)
    cognitive.scan_cognitive([path], tmp_path, True, False, {}, reply_with([GOOD], []))
    seen = []
    assert cognitive.scan_cognitive([path], tmp_path, True, False, {}, reply_with([], seen))
    assert not seen, "unchanged file is not sent twice"
    monkeypatch.setattr(llm, "complete", lambda *a, **k: (_ for _ in ()).throw(AssertionError("no consent, no network")))
    monkeypatch.delenv("GROQ_API_KEY")
    result = scan(tmp_path, ocr=False)
    assert [f["rule"] for f in result["findings"]].count("AI_SENSITIVE") == 1
    d = Document(path)
    d.add_paragraph("changed")
    d.save(path)
    assert "AI_SENSITIVE" not in [f["rule"] for f in scan(tmp_path, ocr=False)["findings"]]


def test_model_switch_rescans_and_failure_keeps_tier1(tmp_path, monkeypatch):
    groq(monkeypatch)
    path = doc(tmp_path)
    cognitive.scan_cognitive([path], tmp_path, True, False, {}, reply_with([GOOD], []))
    monkeypatch.setenv("SHIELD_LLM_MODEL", "other")
    seen = []
    cognitive.scan_cognitive([path], tmp_path, True, False, {}, reply_with([], seen))
    assert seen, "new model, new answer"
    monkeypatch.setattr(llm, "complete", lambda *a, **k: (_ for _ in ()).throw(llm.LLMError("down")))
    (tmp_path / cognitive.CACHE).unlink()
    result = scan(tmp_path, ocr=False, cognitive=True)
    assert "IN_AADHAAR" in [f["rule"] for f in result["findings"]]
    assert "down" in result["stats"]["cognitive"]["error"]


def test_local_only_never_sends(tmp_path, monkeypatch):
    groq(monkeypatch)
    monkeypatch.setenv("SHIELD_LOCAL_ONLY", "1")
    seen, stats = [], {}
    assert not cognitive.scan_cognitive([doc(tmp_path)], tmp_path, True, False, stats, reply_with([GOOD], seen))
    assert not seen and "error" in stats["cognitive"]


def test_image_finding_gets_the_word_box_and_redact_fills_it(tmp_path, monkeypatch):
    groq(monkeypatch)
    from PIL import Image

    Image.new("RGB", (400, 100), "white").save(tmp_path / "shot.png")
    words = [("Patient", (0, 0, 50, 10)), ("record", (60, 0, 50, 10)), ("diagnosis", (0, 20, 60, 10)), ("melanoma", (70, 20, 60, 10)), ("stage", (140, 20, 40, 10)), ("two", (190, 20, 30, 10))]
    monkeypatch.setattr(cognitive, "ocr_words", lambda p: words)
    item = {"value": "melanoma stage two", "kind": "HEALTH", "severity": "high", "reason": "medical condition"}
    found = cognitive.scan_cognitive([tmp_path / "shot.png"], tmp_path, True, True, {}, reply_with([item], []))
    assert found[0].bbox == [70, 20, 150, 10] and "redact_image" in found[0].fixes
    assert cognitive.cached_boxes(tmp_path, "shot.png") == [[70, 20, 150, 10]]
    monkeypatch.setattr("shield.redact.image.ocr_words", lambda p: [])
    monkeypatch.setattr("shield.redact.image._ocr_lines", lambda p: [])
    plan = plan_fix(tmp_path, found[0].to_dict(), "redact_image")
    import base64
    import io

    img = Image.open(io.BytesIO(base64.b64decode(plan.file_ops[0].content_b64)))
    assert img.getpixel((100, 25)) == (0, 0, 0) and img.getpixel((300, 80)) == (255, 255, 255)


def test_unreadable_reply_is_an_error_not_a_clean_file(tmp_path, monkeypatch):
    groq(monkeypatch)
    path = doc(tmp_path)
    empty = httpx.MockTransport(lambda r: httpx.Response(200, json={"choices": [{"message": {"content": ""}}]}))
    stats = {}
    assert cognitive.scan_cognitive([path], tmp_path, True, False, stats, empty) == []
    assert not (tmp_path / cognitive.CACHE).exists()


def test_malformed_items_are_dropped_and_reasons_scrubbed_of_every_value(tmp_path, monkeypatch):
    groq(monkeypatch)
    other = "Ward 4 oncology bed"
    (tmp_path / "n.txt").write_text(f"Clinic list for the day, {DIAGNOSIS}, and {other} assigned.\n")
    junk = [None, "x", {"value": 5}, {"value": DIAGNOSIS, "severity": ["high"], "kind": None, "reason": f"see {other}"},
            {"value": other, "severity": {}, "reason": 7}]
    found = cognitive.scan_cognitive([tmp_path / "n.txt"], tmp_path, True, False, {}, reply_with(junk, []))
    assert len(found) == 2 and all(f.severity == "medium" for f in found)
    assert other not in found[0].rationale


def test_auth_failure_stops_and_keeps_finished_files(tmp_path, monkeypatch):
    groq(monkeypatch)
    for name in ("a.txt", "b.txt", "c.txt"):
        (tmp_path / name).write_text(f"Clinic list {name}: {DIAGNOSIS}, please keep private.\n")
    calls = []

    def handler(request):
        calls.append(1)
        if len(calls) == 1:
            return httpx.Response(200, json={"choices": [{"message": {"content": '{"items":[]}'}}]})
        return httpx.Response(401, json={})

    stats = {}
    files = sorted(tmp_path.glob("*.txt"))
    cognitive.scan_cognitive(files, tmp_path, True, False, stats, httpx.MockTransport(handler))
    assert len(calls) == 2 and "401" in stats["cognitive"]["error"] and stats["cognitive"]["sent"] == 1
    assert list(json.loads((tmp_path / cognitive.CACHE).read_text())) == ["a.txt"], "the next run carries on from b.txt"


def test_damaged_cache_is_a_miss_and_read_only_workspace_keeps_findings(tmp_path, monkeypatch):
    groq(monkeypatch)
    path = doc(tmp_path)
    cache = tmp_path / cognitive.CACHE
    cache.parent.mkdir(parents=True)
    for bad in ("not json", "[1]", '{"note.docx": 3}', '{"note.docx": {"sha": "x"}}'):
        cache.write_text(bad)
        assert cognitive.cached_boxes(tmp_path, "note.docx") == []
        assert cognitive.scan_cognitive([path], tmp_path, False, False, {}) == []
    sha = cognitive._sha(path)
    cache.write_text(json.dumps({"note.docx": {"sha": sha, "model": "m", "findings": [{"old_field": 1}]}}))
    assert cognitive.scan_cognitive([path], tmp_path, False, False, {}) == []
    monkeypatch.setattr(cognitive, "_save", lambda *a: (_ for _ in ()).throw(OSError("read-only")))
    assert cognitive.scan_cognitive([path], tmp_path, True, False, {}, reply_with([GOOD], []))


def test_text_note_gets_a_line_and_a_redact_fix_and_source_code_is_never_sent(tmp_path, monkeypatch):
    groq(monkeypatch)
    from shield.codescan.fix import apply_plan

    (tmp_path / "notes.md").write_text(f"# Clinic\n\nToday: {DIAGNOSIS} and more.\n")
    (tmp_path / "app.py").write_text(f'NOTE = "{DIAGNOSIS} is what this string says, sent never"\n')
    seen = []
    real = cognitive.llm.complete
    monkeypatch.setattr(cognitive.llm, "complete", lambda c, s, u, t=None, **k: real(c, s, u, reply_with([GOOD], seen), **k))
    res = scan(tmp_path, cognitive=True, ocr=False, engines=["pii"])
    assert len(seen) == 1 and "sent never" not in seen[0]
    f = next(x for x in res["findings"] if x["rule"] == "AI_SENSITIVE")
    assert (f["line"], f["col"], f["end_col"]) == (3, 8, 8 + len(DIAGNOSIS))
    apply_plan(plan_fix(tmp_path, f, "redact_pii"), tmp_path)
    assert (tmp_path / "notes.md").read_text() == "# Clinic\n\nToday: [REDACTED-AI] and more.\n"


def test_ignore_hides_a_cached_ai_finding(tmp_path, monkeypatch):
    groq(monkeypatch)
    doc(tmp_path)
    found = cognitive.scan_cognitive([tmp_path / "note.docx"], tmp_path, True, False, {}, reply_with([GOOD], []))
    (tmp_path / ".shieldignore").write_text(found[0].id + "\n")
    assert "AI_SENSITIVE" not in [f["rule"] for f in scan(tmp_path, ocr=False)["findings"]]


def test_strip_exif_keeps_ai_findings_but_redact_clears_them(tmp_path, monkeypatch):
    groq(monkeypatch)
    from PIL import Image

    from shield.codescan.fix import apply_plan

    img = Image.new("RGB", (200, 50), "white")
    exif = Image.Exif()
    exif[0x013B] = "Someone"
    img.save(tmp_path / "s.png", exif=exif)
    words = [("Patient", (0, 0, 50, 10)), ("record", (60, 0, 50, 10)), ("diagnosis", (0, 20, 60, 10)), ("melanoma", (70, 20, 60, 10))]
    monkeypatch.setattr(cognitive, "ocr_words", lambda p: words)
    monkeypatch.setattr("shield.redact.image.ocr_words", lambda p: [])
    monkeypatch.setattr("shield.redact.image._ocr_lines", lambda p: [])
    item = {"value": "melanoma", "kind": "HEALTH", "severity": "high", "reason": "medical"}
    found = cognitive.scan_cognitive([tmp_path / "s.png"], tmp_path, True, True, {}, reply_with([item], []))
    apply_plan(plan_fix(tmp_path, found[0].to_dict(), "strip_exif"), tmp_path)
    assert cognitive.scan_cognitive([tmp_path / "s.png"], tmp_path, False, True, {}), "pixels unchanged"
    apply_plan(plan_fix(tmp_path, found[0].to_dict(), "redact_image"), tmp_path)
    assert not cognitive.scan_cognitive([tmp_path / "s.png"], tmp_path, False, True, {})


def test_claude_reply_with_prose_around_the_json_works(tmp_path, monkeypatch):
    monkeypatch.setenv("SHIELD_LLM_PROVIDER", "anthropic")
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")
    monkeypatch.delenv("SHIELD_LOCAL_ONLY", raising=False)
    monkeypatch.delenv("SHIELD_LLM_MODEL", raising=False)
    path = doc(tmp_path)
    body = "```json\n" + json.dumps({"items": [GOOD]}) + "\n```"
    transport = httpx.MockTransport(lambda r: httpx.Response(200, json={"content": [{"type": "text", "text": body}]}))
    stats = {}
    assert len(cognitive.scan_cognitive([path], tmp_path, True, False, stats, transport)) == 1
    assert stats["cognitive"]["model"] == "anthropic/claude-haiku-5-5"


def test_long_file_is_read_in_chunks_and_a_value_near_the_end_is_found(tmp_path, monkeypatch):
    groq(monkeypatch)
    (tmp_path / "long.log").write_text("".join(f"routine status line {i} ok\n" for i in range(800)) + f"{DIAGNOSIS}\n")
    seen, stats = [], {}
    found = cognitive.scan_cognitive([tmp_path / "long.log"], tmp_path, True, False, stats, reply_with([GOOD], seen))
    assert len(seen) == 2 and [f.line for f in found] == [801] and stats["cognitive"]["truncated"] == 0


def test_repeated_lines_are_sent_once_but_every_copy_is_reported(tmp_path, monkeypatch):
    groq(monkeypatch)
    (tmp_path / "r.log").write_text(f"routine line\n{DIAGNOSIS}\nroutine line\n{DIAGNOSIS}\n" * 50)
    seen = []
    found = cognitive.scan_cognitive([tmp_path / "r.log"], tmp_path, True, False, {}, reply_with([GOOD], seen))
    assert seen[0].count("routine line") == 1 and len(found) == 100
