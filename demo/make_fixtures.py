from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pymupdf
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName
from PIL import Image, ImageDraw

NAME = "Ravi Kumar"
ORG = "Acme Corp"
AADHAAR = "2345 6789 0124"
AADHAAR_BAD = "2345 6789 0120"
PAN = "ABCPK1234F"
MOBILE = "+91 98765 43210"
EMAIL = "ravi.kumar@acme.in"
UPI = "ravi.kumar@okaxis"
IFSC = "SBIN0001234"
ACCOUNT = "123456789012"

VALUES = {
    "NAME": NAME, "ORG": ORG, "AADHAAR": AADHAAR, "AADHAAR_BAD": AADHAAR_BAD,
    "PAN": PAN, "MOBILE": MOBILE, "EMAIL": EMAIL, "UPI": UPI,
    "IFSC": IFSC, "ACCOUNT": ACCOUNT,
}


def _face_image(faces_dir: Path) -> Path:
    src = faces_dir / "victim_1.jpg"
    if src.exists():
        return src
    src.parent.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGB", (300, 300), (230, 220, 210))
    d = ImageDraw.Draw(img)
    d.ellipse((90, 60, 210, 190), fill=(220, 180, 150))
    d.ellipse((120, 110, 135, 125), fill=(20, 20, 20))
    d.ellipse((165, 110, 180, 125), fill=(20, 20, 20))
    d.arc((130, 140, 170, 170), 0, 180, fill=(120, 40, 40), width=4)
    d.rectangle((80, 200, 220, 300), fill=(40, 70, 120))
    img.save(src, "JPEG", quality=90)
    return src


def _pdf(path: Path, face: Path, naive: bool = False):
    doc = pymupdf.open()
    page = doc.new_page()
    text = (
        f"Employee Name: {NAME}\nOrganisation: {ORG}\n"
        f"Aadhaar: {AADHAAR}\nPAN: {PAN}\nMobile: {MOBILE}\n"
        f"Email: {EMAIL}\nUPI: {UPI}\nIFSC: {IFSC}\nAccount: {ACCOUNT}\n"
    )
    page.insert_textbox(pymupdf.Rect(50, 50, 545, 500), text, fontsize=11)
    if face.exists():
        page.insert_image(pymupdf.Rect(360, 50, 560, 250), filename=str(face))
    doc.set_metadata({"author": NAME, "title": "KYC record"})
    try:
        doc.set_xml_metadata(
            "<x:xmpmeta xmlns:x='adobe:ns:meta/'><rdf:RDF xmlns:rdf='http://www.w3.org/1999/02/22-rdf-syntax-ns#'>"
            "<rdf:Description xmlns:dc='http://purl.org/dc/elements/1.1/'><dc:creator>Ravi Kumar</dc:creator>"
            "</rdf:Description></rdf:RDF></x:xmpmeta>"
        )
    except Exception:
        pass
    doc.embfile_add("notes.txt", f"PAN {PAN}".encode())
    annot = page.add_text_annot((72, 600), f"Call {MOBILE}")
    annot.set_info(content=MOBILE)

    if naive:
        rects = page.search_for(AADHAAR)
        for r in rects:
            page.draw_rect(r, color=(0, 0, 0), fill=(0, 0, 0))
    doc.save(path, garbage=4, deflate=True)
    doc.close()


def _add_docx_comment(path: Path, text: str):
    from lxml import etree
    with zipfile.ZipFile(path) as z:
        items = {i.filename: z.read(i.filename) for i in z.infolist()}

    comments = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:comment w:id="0" w:author="Ravi Kumar"><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:comment>
</w:comments>"""
    items["word/comments.xml"] = comments.encode()

    rels_data = items.get("word/_rels/document.xml.rels",
        b"<Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'></Relationships>")
    root = etree.fromstring(rels_data)
    ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    for r in list(root):
        if "comments.xml" in (r.get("Target") or ""):
            root.remove(r)
    el = etree.SubElement(root, f"{ns}Relationship")
    el.set("Id", "rIdComments")
    el.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments")
    el.set("Target", "comments.xml")
    items["word/_rels/document.xml.rels"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    ct = items["[Content_Types].xml"].decode()
    if "comments+xml" not in ct:
        ct = ct.replace(
            "</Types>",
            '<Override PartName="/word/comments.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"/></Types>',
        )
        items["[Content_Types].xml"] = ct.encode()

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in items.items():
            z.writestr(name, data)


def _docx(path: Path):
    doc = Document()
    p = doc.add_paragraph("Aadhaar: ")
    for part in ("2345 ", "6789 ", "0124"):
        p.add_run(part)
    p.add_run(f" PAN: {PAN}")
    p.add_run(f" Mobile: {MOBILE}")

    del_el = OxmlElement("w:del")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = MOBILE
    run.append(t)
    del_el.append(run)
    p._p.append(del_el)

    doc.sections[0].header.paragraphs[0].text = f"PAN {PAN}"
    doc.core_properties.author = NAME
    doc.core_properties.last_modified_by = NAME
    doc.save(path)
    _add_docx_comment(path, EMAIL)


def _xlsx(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = NAME
    ws["B1"] = PAN
    ws["C1"] = f'="UPI:"&"{UPI}"'
    ws["D1"] = IFSC
    ws["D1"].comment = Comment(f"IFSC {IFSC}", NAME)
    ws["E1"] = ACCOUNT

    ws2 = wb.create_sheet("HiddenKyc")
    ws2["A1"] = AADHAAR
    ws2.sheet_state = "veryHidden"

    wb.defined_names["Acct"] = DefinedName("Acct", attr_text="Sheet1!$E$1")
    wb.properties.creator = NAME
    wb.properties.lastModifiedBy = NAME
    wb.save(path)


def _image(path: Path, png: bool = False):
    img = Image.new("RGB", (600, 300), (250, 250, 250))
    d = ImageDraw.Draw(img)
    d.text((20, 40), f"Aadhaar: {AADHAAR}", fill=(0, 0, 0))
    d.text((20, 90), f"PAN: {PAN}", fill=(0, 0, 0))
    if png:
        from PIL.PngImagePlugin import PngInfo
        info = PngInfo()
        info.add_text("Author", NAME)
        img.save(path, "PNG", pnginfo=info)
    else:
        exif = Image.Exif()
        exif[0x013B] = NAME
        img.save(path, "JPEG", exif=exif, quality=90)


def build(out_dir: str | Path) -> dict[str, Path]:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    # faces dir relative to demo/
    faces_dir = Path(__file__).parent / "faces"
    face = _face_image(faces_dir)

    pdf = out / "kyc.pdf"
    naive = out / "naive.pdf"
    docx = out / "kyc.docx"
    xlsx = out / "kyc.xlsx"
    jpg = out / "kyc.jpg"
    png = out / "kyc.png"
    clean = out / "clean.pdf"

    _pdf(pdf, face, naive=False)
    _pdf(naive, face, naive=True)
    _docx(docx)
    _xlsx(xlsx)
    _image(jpg, png=False)
    _image(png, png=True)

    doc2 = pymupdf.open()
    page = doc2.new_page()
    page.insert_textbox(pymupdf.Rect(50, 50, 545, 200), "Quarterly report. No personal data.", fontsize=12)
    doc2.save(clean)
    doc2.close()

    return {
        "pdf": pdf, "naive_pdf": naive, "docx": docx, "xlsx": xlsx,
        "jpg": jpg, "png": png, "clean_pdf": clean, "face": face,
    }


if __name__ == "__main__":
    import sys
    print(build(sys.argv[1] if len(sys.argv) > 1 else "demo/out"))

